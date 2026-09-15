import docx
import os

base = r"D:\workspace\AGI的哲学思考"
files = [
    "1. 关于定义AGI的哲学思考.docx",
    "2. 认知论本体的宿命.docx",
    "3. 热寂、奇点与宇宙的自证.docx",
    "4. 宇宙的自证：当AGI推演撞上儒释道.docx",
]

for f in files:
    path = os.path.join(base, f)
    print("=" * 80)
    print("FILE:", f)
    print("=" * 80)
    doc = docx.Document(path)
    for p in doc.paragraphs:
        if p.text.strip():
            style = p.style.name if p.style else ""
            print(f"[{style}] {p.text}")
    for i, table in enumerate(doc.tables):
        print(f"--- TABLE {i} ---")
        for row in table.rows:
            print(" | ".join(cell.text for cell in row.cells))
    print(f"--- inline_shapes: {len(doc.inline_shapes)} ---")
    print()
