#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""批量将 docx 文件转换为 Markdown 格式"""

import os
from docx import Document

def docx_to_markdown(docx_path):
    """将 docx 文件转换为 Markdown 文本"""
    doc = Document(docx_path)
    md_lines = []
    
    for para in doc.paragraphs:
        text = para.text.strip()
        if not text:
            md_lines.append("")
            continue
        
        # 根据段落样式判断标题级别
        style = para.style.name.lower() if para.style else ""
        
        if "heading 1" in style or "标题 1" in style:
            md_lines.append(f"# {text}")
        elif "heading 2" in style or "标题 2" in style:
            md_lines.append(f"## {text}")
        elif "heading 3" in style or "标题 3" in style:
            md_lines.append(f"### {text}")
        elif "heading 4" in style or "标题 4" in style:
            md_lines.append(f"#### {text}")
        elif "heading 5" in style or "标题 5" in style:
            md_lines.append(f"##### {text}")
        elif "title" in style or "标题" in style:
            md_lines.append(f"# {text}")
        elif "list bullet" in style or "列表" in style:
            md_lines.append(f"- {text}")
        elif "list number" in style:
            md_lines.append(f"1. {text}")
        else:
            # 处理加粗、斜体等行内格式
            line_parts = []
            for run in para.runs:
                run_text = run.text
                if not run_text:
                    continue
                if run.bold and run.italic:
                    line_parts.append(f"***{run_text}***")
                elif run.bold:
                    line_parts.append(f"**{run_text}**")
                elif run.italic:
                    line_parts.append(f"*{run_text}*")
                else:
                    line_parts.append(run_text)
            
            if line_parts:
                md_lines.append("".join(line_parts))
            else:
                md_lines.append(text)
    
    # 处理表格
    for table in doc.tables:
        md_lines.append("")
        # 表头
        header = []
        for cell in table.rows[0].cells:
            header.append(cell.text.strip())
        md_lines.append("| " + " | ".join(header) + " |")
        md_lines.append("| " + " | ".join(["---"] * len(header)) + " |")
        # 表体
        for row in table.rows[1:]:
            row_data = [cell.text.strip() for cell in row.cells]
            md_lines.append("| " + " | ".join(row_data) + " |")
        md_lines.append("")
    
    return "\n".join(md_lines)


def main():
    root_dir = r"D:\workspace\AGI的哲学思考"
    output_dir = os.path.join(root_dir, "哲学起点")
    
    # 找到所有 docx 文件
    docx_files = []
    for f in os.listdir(root_dir):
        if f.lower().endswith(".docx") and not f.startswith("~$"):
            docx_files.append(f)
    
    print(f"Found {len(docx_files)} docx files")
    
    for docx_file in docx_files:
        docx_path = os.path.join(root_dir, docx_file)
        md_file = os.path.splitext(docx_file)[0] + ".md"
        md_path = os.path.join(output_dir, md_file)
        
        try:
            md_content = docx_to_markdown(docx_path)
            with open(md_path, "w", encoding="utf-8") as f:
                f.write(md_content)
            print(f"  Converted: {docx_file} -> {md_file}")
        except Exception as e:
            print(f"  Error converting {docx_file}: {e}")
    
    print(f"\nDone. Output directory: {output_dir}")


if __name__ == "__main__":
    main()
