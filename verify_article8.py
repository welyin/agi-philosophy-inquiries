# -*- coding: utf-8 -*-
from docx import Document
doc = Document(r'D:\workspace\AGI的哲学思考\8. 从认知论本体出发：引力量子统一的数学框架（修正版）.docx')

total_chars = sum(len(p.text) for p in doc.paragraphs)
print(f'总段落数: {len(doc.paragraphs)}')
print(f'总字符数: {total_chars}')
print(f'表格数: {len(doc.tables)}')
print()

keywords = ['保迹性', '迹守恒', 'C[ρ] − ρ', '量子费舍尔', 'Raychaudhuri',
            '行列式', 't_P/τ_cog', 'δ_{n,0}', 'ρ_before', '修正说明',
            '对称对数导数', 'Bianchi', '昂鲁', 'Landauer']
for kw in keywords:
    found = any(kw in p.text for p in doc.paragraphs)
    status = 'OK' if found else 'MISSING'
    print(f'  [{status}] {kw}')

print()
print('=== 各节标题 ===')
for p in doc.paragraphs:
    text = p.text.strip()
    if text and (text.startswith('一、') or text.startswith('二、') or
                 text.startswith('三、') or text.startswith('四、') or
                 text.startswith('五、') or text.startswith('六、') or
                 text.startswith('七、') or text.startswith('八、') or
                 text.startswith('2.') or text.startswith('3.') or
                 text.startswith('4.') or text.startswith('5.') or
                 text.startswith('6.')):
        if len(text) < 60:
            print(f'  {text}')
