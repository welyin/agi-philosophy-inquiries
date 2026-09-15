# -*- coding: utf-8 -*-
from docx import Document
doc = Document(r'D:\workspace\AGI的哲学思考\8. 从认知论本体出发：引力量子统一的数学框架（v2）.docx')

total_chars = sum(len(p.text) for p in doc.paragraphs)
print(f'段落数: {len(doc.paragraphs)}, 总字符: {total_chars}, 表格: {len(doc.tables)}')
print()

# 检查关键推导点
checks = {
    '局部伦德勒度规': 'ds² = −(aξ)²',
    '仿射参数关系': 'λ = −e^{−aη}/a',
    'Raychaudhuri方程': 'dθ/dλ = −R_{μν}',
    '昂鲁温度': 'T_cog = ℏ a / (2π c k_B)',
    '全息熵常数': 'η_H = k_B c³ / (4ℏG)',
    '爱因斯坦方程结果': 'R_{μν} − (1/2) R g_{μν} + Λ g_{μν} = (8πG/c⁴)',
    'Kähler恒等式': 'ω(X, Y) = g(JX, Y)',
    'Fubini-Study度量': 'g = 4( ⟨dψ|dψ⟩',
    '哈密顿矢量场': 'X_A = (i/ℏ)(A − ⟨A⟩)',
    '范数方差关系': '||X_A||² = (4/ℏ²)(ΔA)²',
    '泊松括号对易子': '|{f_A, f_B}| = (2/ℏ²)|⟨[A,B]⟩|',
    'Robertson关系': '(ΔA)(ΔB) ≥ (1/2)|⟨[A,B]⟩|',
    '海森堡结果': 'Δx · Δp ≥ ℏ/2',
    '折叠算子保迹': 'Tr[C[ρ]] = 1',
    '统一方程': '∂ρ/∂t = −(i/ℏ)[H, ρ] + κ · ( C[ρ] − ρ )',
    '迹守恒验证': 'Tr[∂ρ/∂t] =',
    '预言1量纲': '(t_P/τ_cog)²',
    '黑洞谱极限': 'p_n = ⟨E_n|ρ|E_n⟩  →  δ_{n,0}',
}

print('=== 关键推导点检查 ===')
all_ok = True
for name, kw in checks.items():
    found = any(kw in p.text for p in doc.paragraphs)
    status = 'OK' if found else 'MISSING'
    if not found:
        all_ok = False
    print(f'  [{status}] {name}')

print()
# 检查是否还有错误的中间代数
bad_patterns = ['2πc · T_{μν}', 'c³/(4G)) · R', '(8πG/c²) T']
print('=== 错误模式检查 ===')
for bp in bad_patterns:
    found = any(bp in p.text for p in doc.paragraphs)
    label = 'FOUND(BAD)' if found else 'clean'
    print(f'  [{label}] {bp}')

print()
result = '全部通过' if all_ok else '存在缺失'
print(f'总体: {result}')
