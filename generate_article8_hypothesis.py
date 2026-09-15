"""Build the hypothesis manuscript with native Word equations and validate it."""

import json
import re
from copy import deepcopy
from importlib.metadata import version
from pathlib import Path

import pypandoc
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


BASE = Path(__file__).resolve().parent
SOURCE = BASE / "article8_hypothesis.md"
OUTPUT = BASE / "8. 认知论视角下的量子信息与引力涌现（假说稿）.docx"
MARKDOWN_FORMAT = "markdown+tex_math_dollars+fenced_code_blocks"


def walk_nodes(value):
    if isinstance(value, dict):
        yield value
        for child in value.values():
            yield from walk_nodes(child)
    elif isinstance(value, list):
        for child in value:
            yield from walk_nodes(child)


def validate_source():
    text = SOURCE.read_text(encoding="utf-8")
    syntax_tree = json.loads(pypandoc.convert_file(str(SOURCE), "json", format=MARKDOWN_FORMAT))
    nodes = list(walk_nodes(syntax_tree))
    equations = [node for node in nodes if node.get("t") == "Math" and node["c"][0]["t"] == "DisplayMath"]
    equation_numbers = []
    for equation in equations:
        tags = re.findall(r"\\tag\{(\d+)\}", equation["c"][1])
        if len(tags) != 1:
            raise ValueError("Each display equation must have one explicit number.")
        equation_numbers.append(int(tags[0]))
    if equation_numbers != list(range(1, 34)):
        raise ValueError(f"Unexpected equation sequence: {equation_numbers}")
    for number in range(1, 9):
        if not re.search(rf"^# {number} ", text, flags=re.MULTILINE):
            raise ValueError(f"Missing main section {number}.")
    for heading in ("摘要", "附录 A", "附录 B", "研究透明度", "参考文献"):
        if f"# {heading}" not in text:
            raise ValueError(f"Missing required section: {heading}")
    bibliography = text.split("# 参考文献", 1)[1]
    reference_numbers = [int(number) for number in re.findall(r"^\[(\d+)\]", bibliography, flags=re.MULTILINE)]
    if reference_numbers != list(range(1, 9)):
        raise ValueError("Reference numbering must be complete and consecutive.")
    body_nodes = []
    for block in syntax_tree["blocks"]:
        block_nodes = list(walk_nodes(block))
        if block.get("t") == "Header" and any(
            node.get("t") == "Str" and node.get("c") == "参考文献"
            for node in block_nodes
        ):
            break
        body_nodes.extend(block_nodes)
    citations = {
        int(number)
        for node in body_nodes if node.get("t") == "Str"
        for group in re.findall(r"\[(\d+(?:,\d+)*)\]", node["c"])
        for number in group.split(",")
    }
    if citations != set(reference_numbers):
        raise ValueError("Citations and bibliography do not match.")
    if re.search(r"\b(?:TODO|TBD|FIXME)\b|\ufffd", text):
        raise ValueError("The manuscript contains an unresolved placeholder or invalid character.")
    return {
        "source_characters": len(text),
        "display_equations": len(equations),
        "all_math_nodes": sum(node.get("t") == "Math" for node in nodes),
        "tables": sum(node.get("t") == "Table" for node in nodes),
        "references": len(reference_numbers),
    }


def set_style_font(style, east_asia, size, latin="Times New Roman"):
    style.font.name = latin
    style.font.size = Pt(size)
    style.element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), east_asia)


def apply_layout(document):
    section = document.sections[0]
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.3)
    section.bottom_margin = Cm(2.2)
    section.left_margin = Cm(2.4)
    section.right_margin = Cm(2.4)
    section.header_distance = Cm(1)
    section.footer_distance = Cm(1)

    for style_name in ("Normal", "Body Text", "First Paragraph", "Compact"):
        if style_name in document.styles:
            style = document.styles[style_name]
            set_style_font(style, "宋体", 11)
            style.paragraph_format.line_spacing = 1.4
            style.paragraph_format.space_after = Pt(5)
            style.paragraph_format.widow_control = True

    for style_name, size in (("Title", 19), ("Subtitle", 12), ("Heading 1", 15), ("Heading 2", 12)):
        style = document.styles[style_name]
        set_style_font(style, "黑体", size)
        style.font.color.rgb = RGBColor(0x1F, 0x3A, 0x5F)
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.space_before = Pt(14)
        style.paragraph_format.space_after = Pt(7)
    for style_name in ("Title", "Subtitle", "Date"):
        document.styles[style_name].paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_style_font(document.styles["Date"], "宋体", 10)
    if "Source Code" in document.styles:
        set_style_font(document.styles["Source Code"], "宋体", 9, latin="Consolas")

    for number, equation in enumerate(document.element.xpath(".//m:oMathPara/m:oMath"), 1):
        number_run = OxmlElement("m:r")
        math_properties = OxmlElement("m:rPr")
        plain_style = OxmlElement("m:sty")
        plain_style.set(qn("m:val"), "p")
        math_properties.append(plain_style)
        number_run.append(math_properties)
        number_text = OxmlElement("m:t")
        number_text.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
        number_text.text = f"  ({number})"
        number_run.append(number_text)
        equation.append(number_run)

    bibliography_started = False
    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        if text == "参考文献":
            bibliography_started = True
            paragraph.paragraph_format.page_break_before = True
        if paragraph._p.xpath(".//m:oMathPara"):
            paragraph.paragraph_format.keep_together = True
            paragraph.paragraph_format.space_before = Pt(5)
            paragraph.paragraph_format.space_after = Pt(7)
        elif paragraph.style.name in ("Normal", "Body Text", "First Paragraph") and text:
            if bibliography_started and text != "参考文献":
                paragraph.paragraph_format.first_line_indent = Cm(-0.6)
                paragraph.paragraph_format.left_indent = Cm(0.6)
                paragraph.paragraph_format.line_spacing = 1.15
            else:
                paragraph.paragraph_format.first_line_indent = Cm(0.74)
                paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    standard_styles = Document().styles
    for style_name in ("Normal Table", "Light Grid Accent 1"):
        if style_name not in document.styles:
            document.styles.element.append(deepcopy(standard_styles[style_name].element))
    for table in document.tables:
        table.style = "Light Grid Accent 1"
        table.autofit = False
        for column in table.columns:
            column.width = Cm(16.2 / len(table.columns))
        for index, row in enumerate(table.rows):
            row_properties = row._tr.get_or_add_trPr()
            row_properties.append(OxmlElement("w:cantSplit"))
            if index == 0:
                repeat_header = OxmlElement("w:tblHeader")
                repeat_header.set(qn("w:val"), "true")
                row_properties.append(repeat_header)
            for cell in row.cells:
                cell.width = Cm(16.2 / len(table.columns))
                for paragraph in cell.paragraphs:
                    paragraph.paragraph_format.first_line_indent = Cm(0)
                    paragraph.paragraph_format.left_indent = Cm(0)
                    paragraph.paragraph_format.line_spacing = 1.15
                    paragraph.paragraph_format.space_after = Pt(4)
                    for run in paragraph.runs:
                        run.font.size = Pt(9)
                        run.bold = index == 0

    header = section.header.paragraphs[0]
    header.text = "认知论视角下的量子信息与引力涌现 | 假说论文初稿"
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    for run in header.runs:
        run.font.size = Pt(8)
        run.font.color.rgb = RGBColor(0x66, 0x66, 0x66)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    page_field = OxmlElement("w:fldSimple")
    page_field.set(qn("w:instr"), "PAGE")
    footer._p.append(page_field)
    document.core_properties.title = "认知论视角下的量子信息与引力涌现"
    document.core_properties.subject = "理论假说、数学约束与检验路径"
    document.core_properties.author = ""
    document.core_properties.keywords = "认知论, 量子信息, 引力涌现, 量子测量, 可证伪性"


def validate_document(expected):
    document = Document(OUTPUT)
    display_equations = len(document.element.xpath(".//m:oMathPara"))
    all_math = len(document.element.xpath(".//m:oMath"))
    if display_equations != expected["display_equations"]:
        raise ValueError(f"Word display-equation count mismatch: {display_equations}")
    if all_math != expected["all_math_nodes"]:
        raise ValueError(f"Word math-node count mismatch: {all_math}")
    for number, equation in enumerate(document.element.xpath(".//m:oMathPara/m:oMath"), 1):
        equation_text = "".join(element.text or "" for element in equation.iter(qn("m:t")))
        if not equation_text.endswith(f"({number})"):
            raise ValueError(f"The visible number for equation {number} is missing.")
    if len(document.tables) != expected["tables"]:
        raise ValueError("A manuscript table was lost during conversion.")
    titles = [paragraph.text for paragraph in document.paragraphs if paragraph.style.name == "Heading 1"]
    if len(titles) != 13:
        raise ValueError(f"Unexpected top-level heading count: {len(titles)}")
    full_text = "\n".join(document.element.xpath(".//w:t/text() | .//m:t/text()"))
    if "\\begin{" in full_text or "\\tag{" in full_text or "$$" in full_text:
        raise ValueError("Unconverted mathematical markup remains in the Word document.")
    print(json.dumps({
        **expected,
        "word_file": str(OUTPUT),
        "word_bytes": OUTPUT.stat().st_size,
        "word_paragraphs": len(document.paragraphs),
        "word_headings": len(titles),
        "native_word_display_equations": display_equations,
        "native_word_math_nodes": all_math,
        "python_docx_version": version("python-docx"),
        "numpy_version": version("numpy"),
        "pypandoc_binary_version": version("pypandoc-binary"),
        "pandoc_version": str(pypandoc.get_pandoc_version()),
        "validation_scope": "Document structure and conversion, not proof of the physical hypothesis.",
    }, indent=2, ensure_ascii=False))


def main():
    expected = validate_source()
    pypandoc.convert_file(
        str(SOURCE), "docx", format=MARKDOWN_FORMAT,
        outputfile=str(OUTPUT), extra_args=["--standalone"],
    )
    document = Document(OUTPUT)
    apply_layout(document)
    document.save(OUTPUT)
    validate_document(expected)


if __name__ == "__main__":
    main()